from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.db import models
from django.http import HttpResponse
from django.utils import timezone
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from core.models import Cita, MovimientoInsumo, Pago, VentaInsumo
from control_operativo.models import ComisionPago, ComisionVentaInsumo
from core.permissions import PuedeModificarDinero, PuedeVerDinero, PuedeVerFinanzas
from core.serializers.finanzas import PagoSerializer
from core.utilidades import clinica_usuario, primera_clinica, puede_ver_dinero, rol_usuario

def ahora_local():
    now = timezone.now()
    if settings.USE_TZ and timezone.is_aware(now):
        return timezone.localtime(now)
    return now

class PagoViewSet(viewsets.ModelViewSet):
    queryset = (
        Pago.objects.select_related(
            "cita",
            "cita__paciente",
            "cita__servicio",
            "cita__profesional",
        )
        .all()
        .order_by("-fecha_pago", "-id")
    )
    serializer_class = PagoSerializer
    permission_classes = [IsAuthenticated, PuedeVerDinero]

    def get_permissions(self):
        if self.action in {"create", "update", "partial_update", "destroy", "delete_by_cita"}:
            return [IsAuthenticated(), PuedeModificarDinero()]
        return super().get_permissions()

    def get_queryset(self):
        rol = rol_usuario(self.request.user)

        if not puede_ver_dinero(self.request.user):
            return super().get_queryset().none()

        qs = super().get_queryset()

        # El fisioterapeuta puede ver dinero, pero únicamente de sus citas.
        if rol in {"fisioterapeuta", "terapeuta"}:
            qs = qs.filter(cita__profesional=self.request.user)

        cita_id = self.request.query_params.get("cita")
        if cita_id:
            qs = qs.filter(cita_id=cita_id)

        fecha_pago = self.request.query_params.get("fecha_pago")
        desde = self.request.query_params.get("desde")
        hasta = self.request.query_params.get("hasta")

        if fecha_pago:
            qs = qs.filter(fecha_pago=fecha_pago)
        if desde:
            qs = qs.filter(fecha_pago__gte=desde)
        if hasta:
            qs = qs.filter(fecha_pago__lte=hasta)

        profesional = self.request.query_params.get("profesional")
        if rol == "admin" and profesional:
            qs = qs.filter(cita__profesional_id=profesional)

        return qs

    def _recalcular_cita(self, cita):
        pagos_qs = cita.pagos.all()
        total_pagado = (
            pagos_qs.aggregate(total=models.Sum("anticipo")).get("total")
            or Decimal("0")
        )
        pago_base = pagos_qs.order_by("-fecha_pago", "-id").first()

        if pago_base:
            monto_facturado = Decimal(pago_base.monto_facturado or 0)
            descuento_porcentaje = Decimal(
                pago_base.descuento_porcentaje or 0
            )
        else:
            monto_facturado = Decimal(
                cita.monto_final or cita.precio or 0
            )
            descuento_porcentaje = Decimal(
                cita.descuento_porcentaje or 0
            )

        total_con_descuento = max(
            monto_facturado
            - (
                monto_facturado
                * descuento_porcentaje
                / Decimal("100")
            ),
            Decimal("0"),
        )
        restante = max(
            total_con_descuento - total_pagado,
            Decimal("0"),
        )

        cita.descuento_porcentaje = descuento_porcentaje
        cita.monto_final = total_con_descuento
        cita.anticipo = total_pagado
        cita.pagado = restante <= 0 and total_con_descuento > 0
        cita.save(
            update_fields=[
                "descuento_porcentaje",
                "monto_final",
                "anticipo",
                "pagado",
                "actualizado",
            ]
        )

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        cita = getattr(instance, "cita", None)
        instance.delete()

        if cita:
            self._recalcular_cita(cita)

        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(
        detail=True,
        methods=["get"],
        url_path="ticket",
        permission_classes=[IsAuthenticated, PuedeVerDinero],
    )
    def ticket_pdf(self, request, pk=None):
        pago = self.get_object()
        cita = pago.cita
        clinica = primera_clinica()
        rfc_fijo = "MECE000513F74"

        def media_path(filename):
            try:
                return Path(settings.MEDIA_ROOT) / filename
            except Exception:
                return None

        logo_path = media_path("fisionerv.png")
        qr_path = media_path("qr.png")

        def upper(value):
            return (value or "").strip().upper()

        def money(value):
            try:
                return f"$ {Decimal(value or 0):,.2f}".upper()
            except Exception:
                return "$ 0.00"

        def split_chunks(text, width):
            text = upper(text)
            return [
                text[index : index + width]
                for index in range(0, len(text), width)
            ] or [""]

        negocio = (
            (getattr(clinica, "nombre", "") or "").strip()
            or "FISIONERV"
        )
        direccion = (
            (getattr(clinica, "direccion", "") or "").strip()
            or "DIRECCIÓN NO CONFIGURADA"
        )

        paciente = cita.paciente
        paciente_nombre = (
            f"{paciente.nombres} {paciente.apellido_pat} "
            f"{paciente.apellido_mat or ''}"
        ).strip()

        profesional = cita.profesional
        profesional_nombre = (
            f"{profesional.first_name or ''} "
            f"{profesional.last_name or ''}"
        ).strip() or profesional.username

        costo_servicio = Decimal(cita.precio or 0)
        descuento_pct = Decimal(cita.descuento_porcentaje or 0)

        base_descuento = Decimal(
            pago.monto_facturado or costo_servicio
        )
        descuento_monto = (
            base_descuento
            * descuento_pct
            / Decimal("100")
        )

        total_a_pagar = Decimal(cita.monto_final or 0)
        if total_a_pagar <= 0:
            total_a_pagar = max(
                base_descuento - descuento_monto,
                Decimal("0"),
            )

        pagos_qs = cita.pagos.all()
        total_pagado = (
            pagos_qs.aggregate(total=models.Sum("anticipo")).get("total")
            or Decimal("0")
        )
        restante = max(
            total_a_pagar - Decimal(total_pagado),
            Decimal("0"),
        )

        pagos_por_metodo = (
            pagos_qs.values("metodo_pago")
            .annotate(total=models.Sum("anticipo"))
            .order_by("metodo_pago")
        )

        servicio_nombre = (
            (getattr(cita.servicio, "nombre", "") or "").strip()
            or "SERVICIO"
        )

        ticket_width = 80 * mm
        line_height = 4 * mm
        top_padding = 8 * mm
        bottom_padding = 8 * mm

        lineas_estimadas = (
            30
            + len(split_chunks(direccion, 32))
            + len(split_chunks(paciente_nombre, 32))
            + len(split_chunks(profesional_nombre, 32))
            + len(split_chunks(servicio_nombre, 32))
            + max(1, pagos_por_metodo.count())
        )
        altura = (
            top_padding
            + bottom_padding
            + lineas_estimadas * line_height
            + 28 * mm
        )

        buffer = BytesIO()
        pdf = canvas.Canvas(
            buffer,
            pagesize=(ticket_width, altura),
        )
        y = altura - top_padding

        def center(text, font="Helvetica", size=8.5):
            nonlocal y
            pdf.setFont(font, size)
            value = upper(text)
            width = pdf.stringWidth(value, font, size)
            pdf.drawString(
                max(2 * mm, (ticket_width - width) / 2),
                y,
                value,
            )
            y -= line_height

        def left_right(left, right, font="Helvetica", size=8.5):
            nonlocal y
            pdf.setFont(font, size)
            left = upper(left)
            right = upper(right)
            x_left = 4 * mm
            x_right = ticket_width - 4 * mm
            pdf.drawString(x_left, y, left)
            pdf.drawString(
                x_right - pdf.stringWidth(right, font, size),
                y,
                right,
            )
            y -= line_height

        if logo_path and logo_path.exists():
            try:
                image = ImageReader(str(logo_path))
                draw_w = 60 * mm
                draw_h = 18 * mm
                pdf.drawImage(
                    image,
                    (ticket_width - draw_w) / 2,
                    y - draw_h,
                    width=draw_w,
                    height=draw_h,
                    preserveAspectRatio=True,
                    mask="auto",
                )
                y -= draw_h + 3 * mm
            except Exception:
                pass

        separador = "-" * 32
        now = ahora_local()

        center(negocio, "Helvetica-Bold", 9)
        center(f"RFC: {rfc_fijo}")

        for linea in split_chunks(direccion, 32):
            center(linea)

        center(separador)
        center(f"VENTA: {cita.id}  TICKET: {pago.id}")
        center(
            f"EMISIÓN: {now.strftime('%Y-%m-%d %H:%M:%S')}"
        )
        center(f"FECHA PAGO: {pago.fecha_pago}")
        center(separador)

        center("CLIENTE", "Helvetica-Bold")
        for linea in split_chunks(paciente_nombre, 32):
            center(linea)

        center("PROFESIONAL", "Helvetica-Bold")
        for linea in split_chunks(profesional_nombre, 32):
            center(linea)

        center(separador)
        center("DETALLE", "Helvetica-Bold")
        for linea in split_chunks(servicio_nombre, 32):
            center(linea)

        center(separador)
        left_right("COSTO", money(costo_servicio))
        left_right(
            f"DESC ({descuento_pct}%)",
            f"-{money(descuento_monto)}",
        )
        left_right(
            "MONTO FACTURADO",
            money(pago.monto_facturado),
        )

        center(separador)
        left_right("TOTAL A PAGAR", money(total_a_pagar))
        left_right("COBRADO", money(total_pagado))
        left_right("RESTANTE", money(restante))

        center(separador)
        center("PAGOS POR MÉTODO", "Helvetica-Bold")

        metodo_labels = {
            "efectivo": "EFECTIVO",
            "tarjeta": "TARJETA",
            "transferencia": "TRANSFERENCIA",
            "otro": "OTRO",
        }

        if pagos_por_metodo:
            for row in pagos_por_metodo:
                metodo = (row.get("metodo_pago") or "otro").strip()
                left_right(
                    metodo_labels.get(metodo, upper(metodo)),
                    money(row.get("total") or 0),
                )
        else:
            center("SIN PAGOS REGISTRADOS")

        center(separador)
        center("GRACIAS POR SU PREFERENCIA", "Helvetica-Bold")
        center("DOCUMENTO GENERADO POR EL SISTEMA", size=8)

        if qr_path and qr_path.exists():
            try:
                qr_size = 18 * mm
                pdf.drawImage(
                    ImageReader(str(qr_path)),
                    ticket_width - 4 * mm - qr_size,
                    4 * mm,
                    width=qr_size,
                    height=qr_size,
                    preserveAspectRatio=True,
                    mask="auto",
                )
            except Exception:
                pass

        pdf.showPage()
        pdf.save()

        contenido = buffer.getvalue()
        buffer.close()

        filename = (
            f"TICKET_VENTA_{cita.id}_PAGO_{pago.id}.PDF"
        )
        response = HttpResponse(
            contenido,
            content_type="application/pdf",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )
        return response


    def _corte_data(self, request):
        hoy = ahora_local().date()
        desde = request.query_params.get("desde") or request.query_params.get("fecha") or hoy.isoformat()
        hasta = request.query_params.get("hasta") or request.query_params.get("fecha") or desde
        if desde > hasta:
            desde, hasta = hasta, desde

        pagos = self.get_queryset().filter(fecha_pago__range=(desde, hasta))
        profesional = request.query_params.get("profesional")
        if profesional:
            pagos = pagos.filter(cita__profesional_id=profesional)

        clinica = clinica_usuario(request.user) or primera_clinica()
        ventas_insumos = VentaInsumo.objects.select_related("producto", "paciente").filter(fecha_pago__range=(desde, hasta))
        egresos_insumos = MovimientoInsumo.objects.select_related("producto").filter(fecha__range=(desde, hasta), tipo="compra")
        if clinica:
            ventas_insumos = ventas_insumos.filter(producto__clinica=clinica)
            egresos_insumos = egresos_insumos.filter(producto__clinica=clinica)

        registros = []
        por_metodo = {}
        por_servicio = {}
        por_profesional = {}

        for pago in pagos.order_by("fecha_pago", "id"):
            ingreso = Decimal(pago.anticipo or 0)
            paciente = pago.cita.paciente
            profesional_obj = pago.cita.profesional
            cliente = f"{paciente.nombres} {paciente.apellido_pat} {paciente.apellido_mat or ''}".strip()
            profesional_nombre = (profesional_obj.get_full_name() or "").strip() or profesional_obj.username
            servicio = pago.cita.servicio.nombre
            metodo = pago.metodo_pago or "otro"
            por_metodo[metodo] = por_metodo.get(metodo, Decimal("0")) + ingreso
            por_servicio[servicio] = por_servicio.get(servicio, Decimal("0")) + ingreso
            por_profesional[profesional_nombre] = por_profesional.get(profesional_nombre, Decimal("0")) + ingreso
            registros.append({
                "tipo": "servicio", "categoria": "Servicios", "fecha_pago": pago.fecha_pago, "fecha_cita": pago.cita.fecha,
                "nombre_cliente": cliente, "servicio": servicio, "metodo_pago": metodo, "costo": ingreso,
                "egreso_insumo": Decimal("0"), "total_ingreso": ingreso, "total_neto": ingreso,
                "profesional": profesional_nombre, "referencia": pago.comprobante or "", "id": pago.id,
            })

        for venta in ventas_insumos.order_by("fecha_pago", "id"):
            ingreso = Decimal(venta.cantidad or 0) * Decimal(venta.precio_unitario or 0)
            cliente = venta.nombre_cliente or "Mostrador"
            if venta.paciente_id:
                cliente = f"{venta.paciente.nombres} {venta.paciente.apellido_pat} {venta.paciente.apellido_mat or ''}".strip()
            metodo = venta.metodo_pago or "otro"
            por_metodo[metodo] = por_metodo.get(metodo, Decimal("0")) + ingreso
            por_servicio[f"Insumo: {venta.producto.nombre}"] = por_servicio.get(f"Insumo: {venta.producto.nombre}", Decimal("0")) + ingreso
            registros.append({
                "tipo": "insumo", "categoria": "Venta de insumos", "fecha_pago": venta.fecha_pago, "fecha_cita": None,
                "nombre_cliente": cliente, "servicio": venta.producto.nombre, "metodo_pago": metodo, "costo": ingreso,
                "egreso_insumo": Decimal("0"), "total_ingreso": ingreso, "total_neto": ingreso,
                "profesional": "", "referencia": venta.comprobante or "", "id": venta.id,
            })

        egresos = []
        for mov in egresos_insumos.order_by("fecha", "id"):
            monto = Decimal(mov.cantidad or 0) * Decimal(mov.costo_unitario or 0)
            egresos.append({
                "tipo": mov.tipo, "fecha": mov.fecha, "producto": mov.producto.nombre, "cantidad": mov.cantidad,
                "costo_unitario": mov.costo_unitario, "monto": monto, "motivo": mov.motivo or "",
            })
            registros.append({
                "tipo": "egreso", "categoria": "Compra de insumos", "fecha_pago": mov.fecha, "fecha_cita": None,
                "nombre_cliente": mov.motivo or "Proveedor / compra", "servicio": mov.producto.nombre, "metodo_pago": "egreso",
                "costo": Decimal("0"), "egreso_insumo": monto, "total_ingreso": Decimal("0"), "total_neto": -monto,
                "profesional": "", "referencia": mov.motivo or "", "id": mov.id,
            })

        total_servicios = sum((Decimal(r["total_ingreso"]) for r in registros if r["tipo"] == "servicio"), Decimal("0"))
        total_ventas_insumos = sum((Decimal(r["total_ingreso"]) for r in registros if r["tipo"] == "insumo"), Decimal("0"))
        total_ingresos = total_servicios + total_ventas_insumos
        total_egresos = sum((Decimal(item["monto"]) for item in egresos), Decimal("0"))
        # Comisión por pago, con porcentaje histórico congelado al registrarse.
        comisiones = ComisionPago.objects.filter(pago__in=pagos)
        monto_comisiones = comisiones.aggregate(
            mp=models.Sum('monto_mp'), recepcion=models.Sum('monto_recepcion'))
        comisiones_insumos = ComisionVentaInsumo.objects.filter(venta__in=ventas_insumos)
        total_mp_insumos = comisiones_insumos.aggregate(total=models.Sum('monto_mp'))['total'] or Decimal('0')
        comisiones_por_venta = {item.venta_id: item for item in comisiones_insumos}
        total_mp = (monto_comisiones['mp'] or Decimal('0')) + total_mp_insumos
        total_recepcion = monto_comisiones['recepcion'] or Decimal('0')
        comisiones_por_pago = {item.pago_id: item for item in comisiones}
        for registro in registros:
            registro['comision_mp'] = Decimal('0')
            registro['comision_recepcion'] = Decimal('0')
            if registro['tipo'] == 'insumo':
                detalle = comisiones_por_venta.get(registro['id'])
                if detalle:
                    registro['comision_mp'] = detalle.monto_mp
                    registro['total_neto'] -= detalle.monto_mp
            if registro['tipo'] == 'servicio':
                # La ID de movimiento de servicio es la ID del pago.
                detalle = comisiones_por_pago.get(registro['id'])
                if detalle:
                    registro['comision_mp'] = detalle.monto_mp
                    registro['comision_recepcion'] = detalle.monto_recepcion
                    registro['total_neto'] -= detalle.monto_mp + detalle.monto_recepcion
        total_neto = total_ingresos - total_egresos - total_mp - total_recepcion

        return {
            "desde": desde, "hasta": hasta, "criterio_ingresos": "fecha_pago", "registros": registros, "egresos": egresos,
            "total_servicios": total_servicios, "total_ventas_insumos": total_ventas_insumos,
            "total_ingresos": total_ingresos, "total_egresos_insumos": total_egresos,
            "total_comisiones_mp": total_mp, "total_comisiones_recepcion": total_recepcion, "total_neto": total_neto,
            "por_metodo": [{"metodo": k, "total": v} for k, v in sorted(por_metodo.items(), key=lambda x: x[1], reverse=True)],
            "por_servicio": [{"servicio": k, "total": v} for k, v in sorted(por_servicio.items(), key=lambda x: x[1], reverse=True)],
            "por_profesional": [{"profesional": k, "total": v} for k, v in sorted(por_profesional.items(), key=lambda x: x[1], reverse=True)],
        }

    @action(detail=False, methods=["get"], url_path="corte-caja", permission_classes=[IsAuthenticated, PuedeVerFinanzas])
    def corte_caja(self, request):
        return Response(self._corte_data(request))

    @action(detail=False, methods=["get"], url_path="corte-caja-excel", permission_classes=[IsAuthenticated, PuedeVerFinanzas])
    def corte_caja_excel(self, request):
        data = self._corte_data(request)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen"
        navy = "123F98"
        blue = "EAF2FF"
        green = "DCFCE7"
        red = "FEE2E2"
        slate = "F8FAFC"
        white = "FFFFFF"
        thin = Side(style="thin", color="DCE3EC")

        ws.merge_cells("A1:H2")
        ws["A1"] = "FISIONERV · CORTE DE CAJA"
        ws["A1"].font = Font(size=18, bold=True, color=white)
        ws["A1"].fill = PatternFill("solid", fgColor=navy)
        ws["A1"].alignment = Alignment(vertical="center")
        ws["A4"] = "Periodo"
        ws["B4"] = f"{data['desde']} a {data['hasta']}"
        ws["A5"] = "Criterio de ingreso"
        ws["B5"] = "Fecha de pago"
        kpis = [
            ("Ingreso servicios", data["total_servicios"]),
            ("Venta de insumos", data["total_ventas_insumos"]),
            ("Total ingresos", data["total_ingresos"]),
            ("Egresos / insumos", data["total_egresos_insumos"]),
            ("Comisión Mercado Pago", data["total_comisiones_mp"]),
            ("Comisión recepción", data["total_comisiones_recepcion"]),
            ("Total neto", data["total_neto"]),
        ]
        for idx, (label, value) in enumerate(kpis, start=7):
            ws[f"A{idx}"] = label
            ws[f"B{idx}"] = float(value)
            ws[f"A{idx}"].font = Font(bold=True)
            ws[f"A{idx}"].fill = PatternFill("solid", fgColor=slate)
            ws[f"B{idx}"].number_format = '$#,##0.00'
            ws[f"B{idx}"].fill = PatternFill("solid", fgColor=green if label != "Egresos / insumos" else red)
            ws[f"A{idx}"].border = ws[f"B{idx}"].border = Border(bottom=thin)

        start_method = 14
        ws[f"A{start_method}"] = "Método de pago"
        ws[f"B{start_method}"] = "Ingreso"
        for cell in ws[start_method][:2]:
            cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True)
        for i, item in enumerate(data["por_metodo"], start=start_method + 1):
            ws[f"A{i}"] = item["metodo"].title(); ws[f"B{i}"] = float(item["total"]); ws[f"B{i}"].number_format = '$#,##0.00'

        if data["por_metodo"]:
            pie = PieChart(); pie.title = "Ingresos por método de pago"
            labels = Reference(ws, min_col=1, min_row=start_method + 1, max_row=start_method + len(data["por_metodo"]))
            values = Reference(ws, min_col=2, min_row=start_method, max_row=start_method + len(data["por_metodo"]))
            pie.add_data(values, titles_from_data=True); pie.set_categories(labels); pie.height = 8; pie.width = 11; ws.add_chart(pie, "D6")

        ws["D14"] = "Concepto"; ws["E14"] = "Monto"
        for cell in [ws["D14"], ws["E14"]]:
            cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True)
        chart_values = [("Servicios", data["total_servicios"]), ("Venta insumos", data["total_ventas_insumos"]), ("Egresos", data["total_egresos_insumos"]), ("Neto", data["total_neto"])]
        for i, (label, value) in enumerate(chart_values, start=15):
            ws[f"D{i}"] = label; ws[f"E{i}"] = float(value); ws[f"E{i}"].number_format = '$#,##0.00'
        bar = BarChart(); bar.type = "col"; bar.style = 10; bar.title = "Rendimiento del corte"; bar.y_axis.title = "MXN"; bar.x_axis.title = "Concepto"
        values = Reference(ws, min_col=5, min_row=14, max_row=18); cats = Reference(ws, min_col=4, min_row=15, max_row=18)
        bar.add_data(values, titles_from_data=True); bar.set_categories(cats); bar.height = 8; bar.width = 12; ws.add_chart(bar, "D21")

        detail = wb.create_sheet("Detalle corte")
        headers = ["Tipo", "Categoría", "Fecha pago", "Fecha cita", "Nombre de cliente", "Servicio / producto", "Método de pago", "Costo", "Egreso = insumo", "Total ingreso", "Total neto", "Profesional", "Referencia", "Comisión MP", "Comisión recepción"]
        detail.append(headers)
        for cell in detail[1]:
            cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True); cell.alignment = Alignment(horizontal="center")
        for row in data["registros"]:
            detail.append([
                row["tipo"], row["categoria"], row["fecha_pago"], row["fecha_cita"], row["nombre_cliente"], row["servicio"], row["metodo_pago"],
                float(row["costo"]), float(row["egreso_insumo"]), float(row["total_ingreso"]), float(row["total_neto"]), row["profesional"], row["referencia"],
                float(row["comision_mp"]), float(row["comision_recepcion"]),
            ])
        for row in detail.iter_rows(min_row=2, min_col=8, max_col=11):
            for cell in row:
                cell.number_format = '$#,##0.00'
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions

        eg = wb.create_sheet("Egresos insumos")
        eg.append(["Fecha", "Producto", "Cantidad", "Costo unitario", "Total", "Motivo"])
        for cell in eg[1]:
            cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True)
        for item in data["egresos"]:
            eg.append([item["fecha"], item["producto"], float(item["cantidad"]), float(item["costo_unitario"]), float(item["monto"]), item["motivo"]])
        for row in eg.iter_rows(min_row=2, min_col=4, max_col=5):
            for cell in row:
                cell.number_format = '$#,##0.00'
        eg.freeze_panes = "A2"

        for sheet in wb.worksheets:
            for column_cells in sheet.columns:
                letter = get_column_letter(column_cells[0].column)
                width = min(max(max((len(str(cell.value or "")) for cell in column_cells), default=8) + 2, 10), 38)
                sheet.column_dimensions[letter].width = width
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

        output = BytesIO(); wb.save(output); output.seek(0)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="corte_caja_{data["desde"]}_{data["hasta"]}.xlsx"'
        return response

    @action(
        detail=False,
        methods=["delete"],
        url_path=r"by-cita/(?P<cita_id>\d+)",
    )
    def delete_by_cita(self, request, cita_id=None):
        cita = Cita.objects.filter(id=cita_id).first()
        if not cita:
            return Response(status=status.HTTP_204_NO_CONTENT)

        cita.pagos.all().delete()
        self._recalcular_cita(cita)

        return Response(status=status.HTTP_204_NO_CONTENT)


